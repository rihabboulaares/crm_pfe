"""
views.py — Endpoints Django REST Framework pour l'agent CRM MCP.
"""
import asyncio
import logging
import base64
from time import perf_counter

from django.utils import timezone
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser

from superadmin.audit import create_ai_agent_run, create_audit_log, finish_ai_agent_run

logger = logging.getLogger(__name__)


class AgentChatView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        try:
            message = request.data.get("message", "").strip()
            if not message:
                return Response({"error": "Message requis."}, status=400)

            user = request.user
            user_id = f"user_{user.id}"
            username = user.username
            company_id = getattr(user, "company_id", 1) or 1
            company = getattr(user, "company", None)
            started = perf_counter()
            run_log = create_ai_agent_run(
                agent_type="crm",
                company=company,
                launched_by=user,
                query=message,
                status="running",
                started_at=timezone.now(),
            )
            create_audit_log(
                actor=user,
                company=company,
                action="launch_agent",
                module="ai_agents",
                object_id=getattr(run_log, "id", None),
                object_repr="Agent CRM",
                description="Message envoye a l'agent CRM",
                metadata={"message_preview": message[:300]},
            )

            logger.info(
                "CRM agent request user=%s company=%s message_length=%s",
                username,
                company_id,
                len(message),
            )

            from .agent import chat_with_memory

            response = asyncio.run(
                chat_with_memory(
                    user_id=user_id,
                    user_message=message,
                    username=username,
                    company_id=company_id,
                    current_user_id=user.id,
                    current_user_role=getattr(user, "role", ""),
                )
            )

            finish_ai_agent_run(
                run_log,
                status="success",
                messages_generated=1,
                duration_seconds=round(perf_counter() - started, 2),
                metadata={"response_preview": str(response)[:500]},
            )
            logger.info("CRM agent response generated user=%s company=%s", username, company_id)
            return Response({
                "response": response,
                "user": username,
                "user_id": user_id,
            })

        except Exception as e:
            if "run_log" in locals():
                finish_ai_agent_run(
                    run_log,
                    status="failed",
                    error_message=str(e),
                    duration_seconds=round(perf_counter() - started, 2) if "started" in locals() else 0,
                )
            create_audit_log(
                actor=getattr(request, "user", None),
                company=getattr(getattr(request, "user", None), "company", None),
                action="system_error",
                module="ai_agents",
                object_repr="Agent CRM",
                description="Erreur agent CRM",
                metadata={"error": str(e)[:500]},
            )
            logger.exception(
                "CRM agent chat failed user=%s company=%s",
                getattr(getattr(request, "user", None), "username", None),
                getattr(getattr(request, "user", None), "company_id", None),
            )
            return Response({"error": "Impossible de traiter la demande pour le moment."}, status=500)


class AgentFileView(APIView):
    """
    Reçoit un fichier (Excel, PDF, Word) + un message utilisateur.
    Extrait le contenu du fichier et l'envoie à l'agent avec le message.
    """
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]

    def post(self, request):
        try:
            file = request.FILES.get("file")
            message = request.data.get("message", "").strip()

            if not file:
                return Response({"error": "Fichier requis."}, status=400)

            filename = file.name.lower()
            content = file.read()

            # ── Extraction du contenu selon le type de fichier ──
            extracted = self._extract(filename, content)

            if extracted.startswith("❌"):
                return Response({"error": extracted}, status=400)

            # ── Construction du message enrichi ──
            user_message = (
                f"{message}\n\n"
                f"=== FICHIER JOINT : {file.name} ===\n"
                f"{extracted}"
            ) if message else (
                f"Analyse ce fichier et effectue les actions nécessaires.\n\n"
                f"=== FICHIER JOINT : {file.name} ===\n"
                f"{extracted}"
            )

            user = request.user
            user_id = f"user_{user.id}"
            username = user.username
            company_id = getattr(user, "company_id", 1) or 1
            company = getattr(user, "company", None)
            started = perf_counter()
            run_log = create_ai_agent_run(
                agent_type="crm",
                company=company,
                launched_by=user,
                query=f"file={file.name}; message={message[:200]}",
                status="running",
                started_at=timezone.now(),
                metadata={"file": file.name},
            )
            create_audit_log(
                actor=user,
                company=company,
                action="launch_agent",
                module="ai_agents",
                object_id=getattr(run_log, "id", None),
                object_repr="Agent CRM fichier",
                description="Fichier envoye a l'agent CRM",
                metadata={"file": file.name},
            )

            logger.info(
                "CRM agent file request user=%s company=%s filename=%s size=%s",
                username,
                company_id,
                file.name,
                len(content),
            )

            from .agent import chat_with_memory

            response = asyncio.run(
                chat_with_memory(
                    user_id=user_id,
                    user_message=user_message,
                    username=username,
                    company_id=company_id,
                    current_user_id=user.id,
                    current_user_role=getattr(user, "role", ""),
                )
            )

            finish_ai_agent_run(
                run_log,
                status="success",
                messages_generated=1,
                duration_seconds=round(perf_counter() - started, 2),
                metadata={"file": file.name, "response_preview": str(response)[:500]},
            )
            return Response({
                "response": response,
                "file": file.name,
                "user": username,
            })

        except Exception as e:
            if "run_log" in locals():
                finish_ai_agent_run(
                    run_log,
                    status="failed",
                    error_message=str(e),
                    duration_seconds=round(perf_counter() - started, 2) if "started" in locals() else 0,
                )
            create_audit_log(
                actor=getattr(request, "user", None),
                company=getattr(getattr(request, "user", None), "company", None),
                action="system_error",
                module="ai_agents",
                object_repr="Agent CRM fichier",
                description="Erreur agent CRM fichier",
                metadata={"error": str(e)[:500]},
            )
            logger.exception(
                "CRM agent file failed user=%s company=%s filename=%s",
                getattr(getattr(request, "user", None), "username", None),
                getattr(getattr(request, "user", None), "company_id", None),
                getattr(locals().get("file", None), "name", None),
            )
            return Response({"error": "Impossible de traiter le fichier pour le moment."}, status=500)

    def _extract(self, filename: str, content: bytes) -> str:
        """Extrait le texte selon l'extension du fichier."""

        # ── Excel ──────────────────────────────────────────────
        if filename.endswith((".xlsx", ".xls")):
            return self._extract_excel(content, filename)

        # ── PDF ────────────────────────────────────────────────
        elif filename.endswith(".pdf"):
            return self._extract_pdf(content)

        # ── Word ───────────────────────────────────────────────
        elif filename.endswith((".docx", ".doc")):
            return self._extract_word(content)

        # ── CSV ────────────────────────────────────────────────
        elif filename.endswith(".csv"):
            return self._extract_csv(content)

        # ── Texte brut ─────────────────────────────────────────
        elif filename.endswith(".txt"):
            return content.decode("utf-8", errors="ignore")

        else:
            return f"❌ Format non supporté : {filename}. Formats acceptés : xlsx, xls, csv, pdf, docx, doc, txt"

    def _extract_excel(self, content: bytes, filename: str) -> str:
        try:
            import openpyxl
            from io import BytesIO
            wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
            result = []
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                result.append(f"--- Feuille : {sheet_name} ---")
                rows = list(ws.iter_rows(values_only=True))
                if not rows:
                    result.append("(feuille vide)")
                    continue
                # Headers
                headers = [str(h).strip() if h is not None else f"Col{i}" for i, h in enumerate(rows[0])]
                result.append(" | ".join(headers))
                result.append("-" * 60)
                # Data rows
                for row in rows[1:]:
                    if all(v is None for v in row):
                        continue
                    values = [str(v).strip() if v is not None else "" for v in row]
                    result.append(" | ".join(values))
            return "\n".join(result)
        except ImportError:
            return "❌ Module openpyxl manquant. Installe-le : pip install openpyxl"
        except Exception as e:
            return f"❌ Erreur lecture Excel : {str(e)}"

    def _extract_pdf(self, content: bytes) -> str:
        try:
            import pdfplumber
            from io import BytesIO
            result = []
            with pdfplumber.open(BytesIO(content)) as pdf:
                for i, page in enumerate(pdf.pages):
                    result.append(f"--- Page {i + 1} ---")
                    # Tables
                    tables = page.extract_tables()
                    if tables:
                        for table in tables:
                            for row in table:
                                cleaned = [str(cell).strip() if cell else "" for cell in row]
                                result.append(" | ".join(cleaned))
                    else:
                        text = page.extract_text()
                        if text:
                            result.append(text)
            return "\n".join(result) if result else "❌ PDF vide ou illisible"
        except ImportError:
            return "❌ Module pdfplumber manquant. Installe-le : pip install pdfplumber"
        except Exception as e:
            return f"❌ Erreur lecture PDF : {str(e)}"

    def _extract_word(self, content: bytes) -> str:
        try:
            import docx
            from io import BytesIO
            doc = docx.Document(BytesIO(content))
            result = []
            # Paragraphes
            for para in doc.paragraphs:
                if para.text.strip():
                    result.append(para.text.strip())
            # Tables
            for table in doc.tables:
                result.append("--- Tableau ---")
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    result.append(" | ".join(cells))
            return "\n".join(result) if result else "❌ Document vide"
        except ImportError:
            return "❌ Module python-docx manquant. Installe-le : pip install python-docx"
        except Exception as e:
            return f"❌ Erreur lecture Word : {str(e)}"

    def _extract_csv(self, content: bytes) -> str:
        try:
            import csv
            from io import StringIO
            text = content.decode("utf-8", errors="ignore")
            reader = csv.reader(StringIO(text))
            rows = list(reader)
            if not rows:
                return "❌ CSV vide"
            result = []
            headers = rows[0]
            result.append(" | ".join(headers))
            result.append("-" * 60)
            for row in rows[1:]:
                if any(cell.strip() for cell in row):
                    result.append(" | ".join(row))
            return "\n".join(result)
        except Exception as e:
            return f"❌ Erreur lecture CSV : {str(e)}"


class AgentResetView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        from .agent import reset_memory

        user_id = f"user_{request.user.id}"
        result = reset_memory(user_id)
        return Response({"message": result})


class AgentStatusView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        try:
            from .memory import REDIS_AVAILABLE
        except Exception:
            logger.exception("Unable to read CRM agent Redis status")
            REDIS_AVAILABLE = False

        return Response({
            "status": "ok",
            "redis": REDIS_AVAILABLE,
            "mode": "MCP",
        })
